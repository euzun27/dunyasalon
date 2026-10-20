"""Ürün yükleme şablonunu (apps/web/public/urun-sablonu.xlsx) üretir: python3 scripts/make-product-template.py"""
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

PRODUCTS = {
    'Su ve Maden Suyu': ['Su 0,33 L', 'Su 0,5 L', 'Su 1,5 L', 'Maden suyu 200 ml', 'Meyveli maden suyu 200 ml'],
    'Meşrubat': ['Kola 330 ml kutu', 'Kola 1 L', 'Kola 2,5 L', 'Şekersiz kola 330 ml kutu', 'Gazoz 330 ml kutu', 'Portakallı gazoz 330 ml kutu',
                 'Limonlu gazoz 330 ml kutu', 'Soğuk çay şeftali 330 ml', 'Soğuk çay limon 330 ml', 'Enerji içeceği 250 ml'],
    'Ayran ve Süt': ['Ayran 200 ml', 'Ayran 300 ml', 'Ayran 1 L', 'Kefir 200 ml', 'Sütlü içecek 200 ml'],
    'Meyve Suyu': ['Meyve suyu vişne 200 ml', 'Meyve suyu şeftali 200 ml', 'Meyve suyu karışık 200 ml', 'Meyve suyu 1 L', 'Şalgam suyu 300 ml', 'Limonata 250 ml'],
    'Sıcak İçecek': ['Çay (bardak)', 'Kahve (hazır, fincan)', 'Türk kahvesi', 'Sıcak çikolata', 'Ihlamur / bitki çayı'],
    'Atıştırmalık': ['Cips sade (küçük)', 'Cips sade (büyük)', 'Cips baharatlı (küçük)', 'Mısır cipsi', 'Patlamış mısır', 'Kraker', 'Simit / poğaça',
                     'Tuzlu çubuk kraker', 'Bisküvi', 'Gofret', 'Kek'],
    'Çikolata ve Şekerleme': ['Çikolata (tablet, küçük)', 'Çikolata (tablet, büyük)', 'Fındıklı çikolata', 'Çikolatalı bar', 'Sakız', 'Şeker (paket)',
                              'Jelibon (paket)', 'Lolipop', 'Lokum (paket)'],
    'Kuruyemiş': ['Çekirdek (paket)', 'Kabak çekirdeği (paket)', 'Fıstık (paket)', 'Karışık kuruyemiş (paket)', 'Leblebi (paket)'],
    'Dondurma': ['Dondurma külah', 'Dondurma çubuk', 'Dondurma kap'],
    'Hediyelik ve Diğer': ['Kolonya 100 ml', 'Kolonya 400 ml', 'Islak mendil', 'Kağıt mendil (paket)', 'Peçete (paket)', 'Lokum kutusu (hediyelik)',
                           'Çiçek / gelin buketi', 'Yaka çiçeği', 'Takı kurdelesi', 'Pil', 'Telefon şarj kablosu'],
}

wb = Workbook()
ws = wb.active
ws.title = 'Ürünler'
headers = ['Barkod', 'Ürün adı', 'Kategori', 'Satış fiyatı (₺)', 'KDV %', 'Stok (adet)', 'Birim']
widths = [20, 34, 22, 17, 9, 13, 10]
brand = PatternFill('solid', fgColor='7A1F3D')
thin = Side(style='thin', color='E8E1E4')
for i, (h, w) in enumerate(zip(headers, widths), start=1):
    c = ws.cell(row=1, column=i, value=h)
    c.font = Font(bold=True, color='FFFFFF')
    c.fill = brand
    c.alignment = Alignment(vertical='center')
    ws.column_dimensions[c.column_letter].width = w
ws.row_dimensions[1].height = 22
ws.freeze_panes = 'A2'

row = 2
for category, names in PRODUCTS.items():
    for name in names:
        ws.cell(row=row, column=1).number_format = '@'  # barkod metin olarak kalsın (8,69E+12 olmasın)
        ws.cell(row=row, column=2, value=name)
        ws.cell(row=row, column=3, value=category)
        ws.cell(row=row, column=4).number_format = '#,##0.00'
        ws.cell(row=row, column=7, value='adet')
        for col in range(1, 8):
            ws.cell(row=row, column=col).border = Border(bottom=thin)
        row += 1
# Kullanıcının ekleyeceği satırlar için de barkod sütunu metin biçiminde olsun
for r in range(row, 1001):
    ws.cell(row=r, column=1).number_format = '@'
    ws.cell(row=r, column=4).number_format = '#,##0.00'

dv = DataValidation(type='decimal', operator='greaterThanOrEqual', formula1='0', allow_blank=True,
                    error='Fiyat sayı olmalıdır (ör. 25 veya 12,50)', errorTitle='Geçersiz fiyat')
ws.add_data_validation(dv)
dv.add('D2:D1000')

info = wb.create_sheet('Nasıl doldurulur')
info.column_dimensions['A'].width = 110
lines = [
    ('Dünyatek Salon Takip Yazılımı · Ürün yükleme şablonu', True),
    ('', False),
    ('1. "Ürünler" sayfasındaki örnek listeyi kendi kantininize göre düzenleyin: satmadığınız ürünlerin satırını silin, eksikleri en alta ekleyin.', False),
    ('2. "Satış fiyatı" zorunludur. Fiyatı boş olan satır yüklenmez ve hangi satır olduğu ekranda gösterilir.', False),
    ('3. "Barkod" sütununu barkod okuyucuyla doldurun: hücreye tıklayın ve ürünü okutun. Barkodu olmayan ürünler kasada ürün kartına tıklanarak satılır.', False),
    ('4. "KDV %" boş bırakılırsa %10 kabul edilir. Doğru oranı muhasebecinizle kontrol edin.', False),
    ('5. "Stok (adet)" doluysa o miktar stoğa giriş olarak eklenir. Sonradan aynı dosyayı tekrar yüklerseniz stok bir kez daha eklenir; bu yüzden ikinci yüklemede stok sütununu boşaltın.', False),
    ('6. Programda: Ürün & Stok > "Excel\'den yükle" > bu dosyayı seçin. Önce bir önizleme gelir, onaylamadan hiçbir şey kaydedilmez.', False),
    ('', False),
    ('Aynı barkodlu ürün programda zaten varsa yeni ürün açılmaz; adı, kategorisi, fiyatı ve KDV oranı güncellenir.', False),
    ('Sütun başlıklarını değiştirmeyin. Dosyayı .xlsx olarak kaydedin (CSV de kabul edilir).', False),
]
for i, (text, bold) in enumerate(lines, start=1):
    c = info.cell(row=i, column=1, value=text)
    c.alignment = Alignment(wrap_text=True, vertical='top')
    if bold:
        c.font = Font(bold=True, size=13, color='7A1F3D')

wb.save('apps/web/public/urun-sablonu.xlsx')
print(row - 2, 'ürün')
